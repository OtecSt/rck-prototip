# -*- coding: utf-8 -*-
"""Регрессия Разработчик_ВР-76: экспорт в шаблон «Вскрытие резервов» V4.7.2.

Проверки (по миссии ВР-76):
  1. Модуль vr_export: сбор данных из базы прогонов (проект № 2), правила
     вебинара ловят нарушения (ВПП=None → нарушение; срок ≠ «6 месяцев» →
     нарушение; отрицательный ЭЭ → нарушение; доля строкой с «%» → нарушение;
     название не «Оптимизация …» → нарушение; доп. показатели пусты — ок).
  2. zapolnit(): xlsx валиден, пишутся только белые ячейки «Ввод данных»,
     формулы (H27, F30, Списки!B142) остаются формулами — запись НЕ в режиме
     data_only; слайды/валидации/прочие части zip — байт-в-байт из эталона;
     workbook.xml получает fullCalcOnLoad (пересчёт при открытии).
  3. API: /api/vr_check (JSON проверок) и /api/vr_export (xlsx) по проекту № 2.
  4. Контрольные НЕ сломаны: /api/oee (OEE 0,704167), /api/meropriyatiya
     (Σ ΔВПП 23,7), /api/potok_calc (такт 2,00 / ВПП 7004,19).

Запуск: python3 _test_vr76.py (сервер на :8000 должен быть запущен).
"""
import base64
import io
import json
import re
import sqlite3
import sys
import urllib.error
import urllib.request
import zipfile
from pathlib import Path

BASE = "http://localhost:8000"
AUTH = ("рцк", "H8BgJlzfjESRDv4F")
PROT = Path(__file__).resolve().parent

sys.path.insert(0, str(PROT))

результаты = []


def итог(инструмент, статус, детали=""):
    результаты.append((инструмент, статус, детали))
    print(f"[{статус}] {инструмент} {детали}")


def req(method, path, body=None, raw=False):
    r = urllib.request.Request(BASE + path, method=method)
    r.add_header("Authorization", "Basic " + base64.b64encode(
        f"{AUTH[0]}:{AUTH[1]}".encode()).decode())
    data = None
    if body is not None:
        data = json.dumps(body, ensure_ascii=False).encode()
        r.add_header("Content-Type", "application/json")
    try:
        with urllib.request.urlopen(r, data=data, timeout=60) as resp:
            blob = resp.read()
            return blob if raw else json.loads(blob.decode())
    except urllib.error.HTTPError as e:
        return {"ok": False, "errors": [f"HTTP {e.code}: {e.read()[:300]}"]}


# --- 1. Модуль: сбор данных из базы (проект № 2 — ОКЗ) ---
import vr_export  # noqa: E402

con = sqlite3.connect(PROT / "ui" / "agent_ceh.db")
con.row_factory = sqlite3.Row
data = vr_export.sobrat_dannye(2, con)
con.close()

ок = (data["предприятие"] == "АО «ОКЗ»"
      and data["срок_реализации"] == "6 месяцев"
      and data["доп_выручка_млн"] == 14 and data["доп_прибыль_млн"] == 8
      and data["выработка_до"] == 3.9 and data["выработка_после"] == 6.0
      and data["впп_до"] is not None and data["впп_после"] is not None
      and data["впп_ед"] == "мин"
      and data["численность_до"] == 10
      and data["доля_выручки"] is None and not data["проблемы"]
      and data["предупреждения"])
итог("vr_export.sobrat_dannye (проект № 2)", "PASS" if ок else "FAIL",
     f"ЭЭ 14/8 млн, выработка 3,9→6,0, ВПП {data['впп_до']}→{data['впп_после']} "
     f"мин, численность {data['численность_до']}; предупреждений "
     f"{len(data['предупреждения'])} (пустые поля честно None, не выдуманы)"
     if ок else json.dumps(data, ensure_ascii=False, default=str)[:400])

# --- 1б. Правила вебинара ловят нарушения ---
def _по_правилу(проверки, имя):
    return next((п for п in проверки if п["правило"] == имя), None)

# эталон нарушения: ВПП=None
без_впп = dict(data, впп_до=None, впп_после=None)
ст = _по_правилу(vr_export.proverit_pravila(без_впп), "ВПП обязателен")
ок1 = ст and ст["статус"] == "нарушение"

# срок ≠ 6 мес
плохой_срок = dict(data, срок_реализации="12 месяцев")
ст = _по_правилу(vr_export.proverit_pravila(плохой_срок), "Срок реализации")
ок2 = ст and ст["статус"] == "нарушение"

# отрицательный ЭЭ → нарушение (правило нуля)
отриц = dict(data, доп_прибыль_млн=-5)
ст = _по_правилу(vr_export.proverit_pravila(отриц), "Дополнительная прибыль")
ок3 = ст and ст["статус"] == "нарушение"

# доля строкой с %
доля_стр = dict(data, доля_выручки="25%")
ст = _по_правилу(vr_export.proverit_pravila(доля_стр), "Доля в выручке")
ок4 = ст and ст["статус"] == "нарушение"

# доля числом — ок
доля_ок = dict(data, доля_выручки=25)
ст = _по_правилу(vr_export.proverit_pravila(доля_ок), "Доля в выручке")
ок5 = ст and ст["статус"] == "ок"

# название не по структуре (у проекта № 2 «ОКЗ — экзамен…»)
ст = _по_правилу(vr_export.proverit_pravila(data), "Название проекта")
ок6 = ст and ст["статус"] == "нарушение"
назв_ок = dict(data, название_проекта="Оптимизация производства комбикорма ПК 1-1")
ст = _по_правилу(vr_export.proverit_pravila(назв_ок), "Название проекта")
ок7 = ст and ст["статус"] == "ок"

# доп. показатели пусты (не в протоколе) — ок; из протокола — ок
ст = _по_правилу(vr_export.proverit_pravila(data), "Дополнительные показатели")
ок8 = ст and ст["статус"] == "ок"

итог("vr_export.proverit_pravila (правила вебинара)",
     "PASS" if all([ок1, ок2, ок3, ок4, ок5, ок6, ок7, ок8]) else "FAIL",
     f"ВПП=None→нарушение {ок1}; срок 12 мес→нарушение {ок2}; ЭЭ −5→нарушение "
     f"{ок3}; доля «25%»→нарушение {ок4}; доля 25→ок {ок5}; название «ОКЗ…»→"
     f"нарушение {ок6}, «Оптимизация…»→ок {ок7}; доп. показатели пусты→ок {ок8}")

# --- 2. Заполнение: валидность и сохранность формул/частей ---
байты = vr_export.zapolnit(data)
z = zipfile.ZipFile(io.BytesIO(байты))
s1 = z.read("xl/worksheets/sheet1.xml").decode("utf-8")
wbx = z.read("xl/workbook.xml").decode("utf-8")
эталон = zipfile.ZipFile(PROT / "assets" / "vr_shablon_v472.xlsx")
прочие_целы = all(эталон.read(n) == z.read(n) for n in эталон.namelist()
                if n not in ("xl/worksheets/sheet1.xml", "xl/workbook.xml"))
формулы_целы = ('<c r="H27" s="74" t="e"><f>(F27-E27)/E27</f>' in s1
                and "<f>ROUNDUP((E29*H27)+(E29-F29),0)</f>" in s1
                and "x14:dataValidations" in s1)
import openpyxl  # noqa: E402
wb = openpyxl.load_workbook(io.BytesIO(байты))  # валиден, читается
ws = wb["Ввод данных"]
значения_ок = (ws["D4"].value == "РЦК" and ws["D5"].value == "АО «ОКЗ»"
               and ws["D11"].value == "6 месяцев" and ws["D23"].value == 14
               and ws["D24"].value == 8 and ws["E28"].value is not None
               and ws["G28"].value == "мин")
ок = прочие_целы and формулы_целы and значения_ок \
    and 'fullCalcOnLoad="1"' in wbx
итог("vr_export.zapolnit (целостность xlsx)", "PASS" if ок else "FAIL",
     f"прочие части zip байт-в-байт: {прочие_целы}; формулы и выпадающие "
     f"списки на месте: {формулы_целы}; значения D4/D5/D11/D23/D24/E28/G28: "
     f"{значения_ок}; fullCalcOnLoad: {'fullCalcOnLoad=\"1\"' in wbx}")

# --- 3. API ---
res = req("POST", "/api/vr_check", {"id_proekta": 2})
проверки = res.get("проверки") or []
ок = (res.get("ok") and проверки
      and _по_правилу(проверки, "Срок реализации")["статус"] == "ок"
      and _по_правилу(проверки, "ВПП обязателен")["статус"] == "ок"
      and _по_правилу(проверки, "Название проекта")["статус"] == "нарушение"
      and _по_правилу(проверки, "Доля в выручке")["статус"] == "нет_данных"
      and res.get("предупреждения"))
итог("/api/vr_check (проект № 2)", "PASS" if ок else "FAIL",
     f"проверок {len(проверки)}, нарушений {res.get('нарушений')}, нет данных "
     f"{res.get('нет_данных')}; название проекта № 2 честно помечено "
     f"нарушением структуры" if ок else json.dumps(res, ensure_ascii=False)[:300])

res = req("POST", "/api/vr_check", {"id_proekta": 999999})
ок_404 = res.get("ok") is False and res.get("errors")
итог("/api/vr_check (несуществующий проект → ошибка, не заглушка)",
     "PASS" if ок_404 else "FAIL", (res.get("errors") or [""])[0][:80])

blob = req("POST", "/api/vr_export", {"id_proekta": 2}, raw=True)
ок = (isinstance(blob, bytes) and blob[:2] == b"PK"
      and zipfile.ZipFile(io.BytesIO(blob)).read(
          "xl/worksheets/sheet1.xml").decode("utf-8").find("АО «ОКЗ»") > 0)
итог("/api/vr_export (проект № 2 → xlsx)", "PASS" if ок else "FAIL",
     f"получено {len(blob) if isinstance(blob, bytes) else blob} байт, "
     "валидный zip, предприятие записано в «Ввод данных»")

# --- 4. Контрольные НЕ сломаны ---
# /api/oee (эталон OEE 0,704167)
res = req("POST", "/api/oee", {
    "поток": "Линия (контроль ВР-76)", "плановое_время_мин": 480,
    "простои": {"аварии_мин": 47, "переналадка_мин": 38,
                "микроостановки_мин": 22},
    "идеальное_тц_мин": 1.0, "выпуск_всего_шт": 350, "брак_шт": 12})
ит = res.get("итоги") or {}
ок = res.get("ok") and ит.get("OEE") is not None \
    and abs(ит["OEE"] - 0.704167) < 1e-4
итог("/api/oee (контроль)", "PASS" if ок else "FAIL",
     f"OEE {ит.get('OEE')} (эталон 0,704167)")

# /api/meropriyatiya (Σ ΔВПП 23,7)
from _test_meropriyatiya_l2 import ЭТАЛОН_ОКЗ  # noqa: E402

res = req("POST", "/api/meropriyatiya", ЭТАЛОН_ОКЗ)
ит2 = res.get("итоги") or {}
ок = (res.get("ok") and ит2.get("сумма_dVPP_pp") is not None
      and abs(ит2["сумма_dVPP_pp"] - 23.7) < 0.11
      and ит2["сумма_dVPP_pp"] <= 23.7 + 1e-9)
итог("/api/meropriyatiya (контроль)", "PASS" if ок else "FAIL",
     f"Σ ΔВПП {ит2.get('сумма_dVPP_pp')} п.п. (эталон 23,7)")

# /api/potok_calc (такт 2,00 / ВПП 7004,19)
res = req("POST", "/api/potok_calc", {
    "поток": "Тестовый поток", "спрос_шт_в_период": 240,
    "доступное_время_мин": 480, "численность_чел": 5,
    "операции": [{"имя": "Оп1", "тц_мин": 1.25, "запас_после_шт": 2000},
                 {"имя": "Оп2", "тц_мин": 1.5, "запас_после_шт": 700},
                 {"имя": "Оп3", "тц_мин": 1.44, "запас_после_шт": 800}]})
ит3 = res.get("итоги") or {}
ок = (ит3.get("такт_мин") == 2.0 and ит3.get("впп_мин") is not None
      and abs(ит3["впп_мин"] - 7004.19) < 0.01)
итог("/api/potok_calc (контроль)", "PASS" if ок else "FAIL",
     f"такт {ит3.get('такт_мин')} (эталон 2,00), ВПП {ит3.get('впп_мин')} "
     f"(эталон 7004,19)")

# --- Итог ---
провалы = [x for x in результаты if x[1] != "PASS"]
print(f"\nИТОГ: {len(результаты) - len(провалы)}/{len(результаты)} PASS")
sys.exit(1 if провалы else 0)
