# -*- coding: utf-8 -*-
import openpyxl
F = "/Users/aleksandr/Desktop/Методички/ИИ-направление/Прототип/разведка_ээ/1.4 Справка плановый_ЭЭ_Оренбургхладокомбинат.xlsm"
wb = openpyxl.load_workbook(F, data_only=False)
wbv = openpyxl.load_workbook(F, data_only=True)

def fmt(v):
    if v is None: return ""
    if isinstance(v, float): return f"{v!r}"
    return str(v).replace("\n"," ⏎ ")[:400]

for sn in ['Чек-лист','Экономический эффект']:
    ws, wsv = wb[sn], wbv[sn]
    out = open(f"/Users/aleksandr/Desktop/Методички/ИИ-направление/Прототип/разведка_ээ/_tmp/dump_{sn}.txt","w")
    out.write(f"SHEET {sn} state={ws.sheet_state}\n")
    for row in ws.iter_rows():
        for c in row:
            v = c.value; cv = wsv[c.coordinate].value
            if v is None and cv is None: continue
            if isinstance(v,str) and v.startswith("="):
                out.write(f"{c.coordinate}\tF: {fmt(v)}\tV: {fmt(cv)}\n")
            elif v is None and cv is not None:
                out.write(f"{c.coordinate}\t(only cached) V: {fmt(cv)}\n")
            else:
                out.write(f"{c.coordinate}\t= {fmt(v)}\n")
    out.close()
    print(sn, "dumped")
