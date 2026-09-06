# -*- coding: utf-8 -*-
import openpyxl
F = "/Users/aleksandr/Desktop/Методички/ИИ-направление/Прототип/разведка_ээ/1.4 Справка плановый_ЭЭ_Оренбургхладокомбинат.xlsm"
wb = openpyxl.load_workbook(F, data_only=False)
wbv = openpyxl.load_workbook(F, data_only=True)

def fmt(v):
    if v is None: return ""
    if isinstance(v, float): return f"{v!r}"
    return str(v).replace("\n"," ⏎ ")

for sn in ['Tech','Изменения','списки','Инструкция','Сверхурочные','Реестр моделей ЭЭ']:
    ws, wsv = wb[sn], wbv[sn]
    print(f"\n########## SHEET {sn} (state={ws.sheet_state}) ##########")
    for row in ws.iter_rows():
        for c in row:
            v = c.value; cv = wsv[c.coordinate].value
            if v is None and cv is None: continue
            s = f"{c.coordinate}: F={fmt(v)}"
            if isinstance(v,str) and v.startswith("="):
                s += f" | V={fmt(cv)}"
            print(s)

# search etalon numbers everywhere
print("\n########## ETALON SEARCH ##########")
targets = [8350501.24, 14299652.41]
for ws in wbv.worksheets:
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if isinstance(v,(int,float)):
                for t in targets:
                    if abs(v-t) < 1.0:
                        print(f"{ws.title}!{c.coordinate} = {v!r} (~{t})  formula={wb[ws.title][c.coordinate].value}")
