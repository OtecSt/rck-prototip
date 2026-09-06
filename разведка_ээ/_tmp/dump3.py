# -*- coding: utf-8 -*-
import openpyxl, zipfile
F = "/Users/aleksandr/Desktop/Методички/ИИ-направление/Прототип/разведка_ээ/1.4 Справка плановый_ЭЭ_Оренбургхладокомбинат.xlsm"
wb = openpyxl.load_workbook(F, data_only=False)
wbv = openpyxl.load_workbook(F, data_only=True)

def fmt(v):
    if v is None: return ""
    if isinstance(v, float): return repr(v)
    return str(v).replace("\n"," ⏎ ")[:300]

for sn in ['Пример расчета','Декомпозиция ЭЭ']:
    ws, wsv = wb[sn], wbv[sn]
    out = open(f"/Users/aleksandr/Desktop/Методички/ИИ-направление/Прототип/разведка_ээ/_tmp/dump_{sn}.txt","w")
    for row in ws.iter_rows():
        for c in row:
            v = c.value; cv = wsv[c.coordinate].value
            if v is None and cv is None: continue
            if hasattr(v,'text'): v = f"ARRAY:{v.text} ref={v.ref}"
            if isinstance(v,str) and v.startswith("="):
                out.write(f"{c.coordinate}\tF: {fmt(v)}\tV: {fmt(cv)}\n")
            elif v is None:
                out.write(f"{c.coordinate}\t(cached) V: {fmt(cv)}\n")
            else:
                out.write(f"{c.coordinate}\t= {fmt(v)}\n")
    out.close(); print(sn,"ok")

# Tech array formulas text
ws = wb['Tech']
print("\n=== TECH ARRAY FORMULAS ===")
for row in ws.iter_rows():
    for c in row:
        v = c.value
        if hasattr(v,'text'): print(c.coordinate, v.ref, v.text)

# data validations
print("\n=== DATA VALIDATIONS ===")
for sn in wb.sheetnames:
    ws = wb[sn]
    for dv in ws.data_validations.dataValidation:
        print(sn, dv.sqref, dv.type, dv.formula1, dv.formula2)

# protection details
print("\n=== PROTECTION ===")
for sn in wb.sheetnames:
    p = wb[sn].protection
    if p.sheet: print(sn, "password_hash=", p.password, "formatCells=",p.formatCells)

# external links
z = zipfile.ZipFile(F)
for n in ['xl/externalLinks/externalLink1.xml','xl/externalLinks/_rels/externalLink1.xml.rels','xl/externalLinks/externalLink2.xml','xl/externalLinks/_rels/externalLink2.xml.rels','docProps/core.xml','docProps/custom.xml','docProps/app.xml']:
    print("\n===",n,"===")
    print(z.read(n).decode('utf-8',errors='replace')[:2000])
