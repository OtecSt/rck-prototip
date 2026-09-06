# -*- coding: utf-8 -*-
import openpyxl, json, zipfile, os
F = "/Users/aleksandr/Desktop/Методички/ИИ-направление/Прототип/разведка_ээ/1.4 Справка плановый_ЭЭ_Оренбургхладокомбинат.xlsm"

# raw zip inventory
z = zipfile.ZipFile(F)
print("=== ZIP INVENTORY ===")
for n in z.namelist():
    print(n, z.getinfo(n).file_size)

wb = openpyxl.load_workbook(F, keep_vba=False, data_only=False)
wbv = openpyxl.load_workbook(F, keep_vba=False, data_only=True)

print("\n=== SHEETS ===")
for ws in wb.worksheets:
    print(f"{ws.title!r} state={ws.sheet_state} dims={ws.dimensions} max_row={ws.max_row} max_col={ws.max_column} protection={ws.protection.sheet}")

print("\n=== DEFINED NAMES ===")
for name, dn in wb.defined_names.items():
    print(name, "->", dn.value, "hidden=", dn.hidden)

print("\n=== WORKBOOK PROTECTION ===", wb.security)

print("\n=== HIDDEN ROWS/COLS per sheet ===")
for ws in wb.worksheets:
    hr = [r for r,d in ws.row_dimensions.items() if d.hidden]
    hc = [c for c,d in ws.column_dimensions.items() if d.hidden]
    if hr or hc:
        print(ws.title, "hidden_rows:", hr, "hidden_cols:", hc)
