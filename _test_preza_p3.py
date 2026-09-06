# -*- coding: utf-8 -*-
"""Регрессия П3 (Спека П1 «Фаст-презентация»): PPTX-движок preza_pptx
и POST /api/preza/pptx поверх сборщика П2.

Проверки (по миссии П3):
  1. Сборка на фикстуре П2 (все инструменты) → 200, content-type pptx,
     файл ненулевой;
  2. Round-trip валидация: python-pptx открывает, 16:9 (13,333×7,5"),
     слайдов = собранным (8), все run'ы шрифтом Calibri;
  3. Эталонные числа в тексте слайдов символ в символ: такт 2,00 и
     ВПП 7 000,00 (симуляция), УМ «фасовка» 38,68 % (как хранит прогон),
     ЭЭ 23 757 960,67, OEE 0,704167;
  4. Speaker notes каждого слайда с источниками содержат «прогон №»
     (у титула — «карточка проекта»);
  5. Дисклеймер «Модельный прогноз, не замер.» на слайде симуляции;
  6. Тема «графит» меняет фон контент-слайдов на #2B2620;
  7. Чек-лист слайдов: два запрошенных → ровно два слайда в файле;
  8. Визуальная приёмка: рендер слайдов в PNG через soffice (если
     установлен) в ui/_qa_p3_slides/, каждый файл ненулевой.

Запуск: python3 _test_preza_p3.py (сервер НЕ нужен — Flask test_client,
временная БД в tempdir).
"""
import copy
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

PROT = Path(__file__).resolve().parent
UI = PROT / "ui"
sys.path.insert(0, str(UI))
sys.path.insert(0, str(PROT))

import app as flask_app  # noqa: E402
import preza_pptx  # noqa: E402
from pptx import Presentation  # noqa: E402

результаты = []


def итог(проверка, статус, детали=""):
    результаты.append((проверка, статус, детали))
    print(f"[{статус}] {проверка} {детали}")


# --- Изоляция: временная БД
_tmp = Path(tempfile.mkdtemp(prefix="preza_p3_")) / "test_agent_ceh.db"
_tmp.touch()
flask_app.DB_PATH = _tmp
flask_app.init_db()
client = flask_app.app.test_client()

ЭТАЛОН = {"поток": "Тестовый поток", "спрос_шт_в_период": 240,
          "доступное_время_мин": 480, "численность_чел": 5,
          "операции": [
              {"имя": "Оп1", "тц_мин": 1.25, "запас_после_шт": 2000},
              {"имя": "Оп2", "тц_мин": 1.5, "запас_после_шт": 700},
              {"имя": "Оп3", "тц_мин": 1.44, "запас_после_шт": 800}]}

# Пара «до/после» симуляции: «до» — эталон нулевой вариативности
# (такт 2,00 / ВПП 7000,00), «после» — Тц УМ 1,5 → 1,2.
СИМ_ДО = dict(ЭТАЛОН)
СИМ_ДО["симуляция"] = {"seed": 1, "вариативность_мин": 1.0,
                       "вариативность_макс": 1.0}
СИМ_ПОСЛЕ = copy.deepcopy(СИМ_ДО)
СИМ_ПОСЛЕ["операции"][1]["тц_мин"] = 1.2

ЭТАЛОН_OEE = {"поток": "Тестовый поток", "плановое_время_мин": 480,
              "простои": {"аварии_мин": 47, "переналадка_мин": 38,
                          "микроостановки_мин": 22},
              "идеальное_тц_мин": 1.0, "выпуск_всего_шт": 350, "брак_шт": 12}

ЭТАЛОН_МЕР = {"predpriyatie": "Тест-завод", "uchastok": "фасовка",
              "vpp_pct": 38.7, "norma_vpp_pct": 15.0,
              "razryv": {"value": 23.7, "unit": "п.п. ВПП", "evidence": "fact"},
              "priznaki": {"простои_оборудования": True,
                           "организация_рабочих_мест": True},
              "vidy_potery": [{"nazvanie": "Ожидание"}], "zamery": {}}

ЭТАЛОН_SMART = {"рамка": "федеральная", "предприятие": "Тест-завод",
                "поток": "Тестовый поток",
                "цели": [{"наименование": "Выработка", "ед": "шт/чел",
                          "текущий": 10, "целевой": 12, "срок": "2026-12-31"}]}


def сохранить(pid, инструмент, вход, метка=None):
    body = {"id_proekta": pid, "инструмент": инструмент, "вход": вход}
    if метка:
        body["метка"] = метка
    j = json.loads(client.post("/api/progony",
                               data=json.dumps(body, ensure_ascii=False),
                               content_type="application/json").data)
    assert j.get("ok"), f"прогон {инструмент} не сохранился: {j}"
    return j["id"]


pid = json.loads(client.post("/api/proekty", json={
    "название": "Тест презы П3", "предприятие": "АО «Тест-завод»",
    "поток": "Тестовый поток"}).data)["id"]
сохранить(pid, "potok_calc", ЭТАЛОН)
сохранить(pid, "uzkie_mesta",
          json.load(open(PROT / "данные_окз_слепой.json", encoding="utf-8")))
сохранить(pid, "simulation", СИМ_ДО, метка="до")
сохранить(pid, "simulation", СИМ_ПОСЛЕ, метка="после")
сохранить(pid, "ee",
          json.load(open(PROT / "данные_ээ_новохром.json", encoding="utf-8")))
сохранить(pid, "oee", ЭТАЛОН_OEE)
сохранить(pid, "meropriyatiya", ЭТАЛОН_МЕР)
сохранить(pid, "smart", ЭТАЛОН_SMART)
сохранить(pid, "protokol",
          json.load(open(PROT / "данные_протокол_кушкуль.json", encoding="utf-8")))

# === 1. Сборка PPTX через API ===============================================
OUT_DIR = UI / "_qa_p3_slides"
OUT_DIR.mkdir(exist_ok=True)
PPTX_ПУТЬ = OUT_DIR / "Защита_тестовый_поток.pptx"
r = client.post("/api/preza/pptx", data=json.dumps({"проект_id": pid}),
                content_type="application/json")
ct = r.headers.get("Content-Type", "")
ок = (r.status_code == 200 and "presentationml.presentation" in ct
      and len(r.data) > 10000)
итог("API: POST /api/preza/pptx → 200, pptx, файл ненулевой",
     "PASS" if ок else "FAIL",
     f"HTTP {r.status_code}, {ct.split(';')[0]}, {len(r.data)} байт")
PPTX_ПУТЬ.write_bytes(r.data)
имя = r.headers.get("Content-Disposition", "")
from urllib.parse import unquote  # noqa: E402
имя_раскр = unquote(имя)
итог("API: имя файла «Защита_<поток>.pptx»",
     "PASS" if "Защита_" in имя_раскр and ".pptx" in имя_раскр else "FAIL",
     имя_раскр[:100])

# === 2. Round-trip валидация ================================================
ошибки_вал = preza_pptx.проверить_pptx(PPTX_ПУТЬ, 8)
итог("round-trip: 16:9, 8 слайдов, шрифт Calibri",
     "PASS" if not ошибки_вал else "FAIL", "; ".join(ошибки_вал) or "ок")

prs = Presentation(str(PPTX_ПУТЬ))
тексты_слайдов = []
for s in prs.slides:
    t = []
    for shp in s.shapes:
        if shp.has_text_frame:
            t.append(shp.text_frame.text)
    тексты_слайдов.append("\n".join(t))

# === 3. Эталонные числа символ в символ =====================================
весь_текст = "\n".join(тексты_слайдов)
эталоны = {
    "такт 2,00": "2,00" in тексты_слайдов[2],
    "ВПП 7 000,00 (симуляция «до»)": "7 000,00" in тексты_слайдов[3],
    "УМ фасовка 38,68 %": ("фасовк" in тексты_слайдов[1]
                           and "38,68 %" in тексты_слайдов[1]),
    "ЭЭ 23 757 960,67": "23 757 960,67" in тексты_слайдов[5],
    "OEE 0,704167": "0,704167" in тексты_слайдов[2],
}
for имя, ок_э in эталоны.items():
    итог(f"эталон на слайде: {имя}", "PASS" if ок_э else "FAIL")

# === 4. Speaker notes ========================================================
заметки = [s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
           for s in prs.slides]
ок_notes = all("прогон №" in з for з in заметки[1:]) \
    and "карточка проекта" in заметки[0]
итог("speaker notes: «источник: прогон № N» на слайдах с данными",
     "PASS" if ок_notes else "FAIL",
     f"слайдов с «прогон №»: {sum('прогон №' in з for з in заметки)}/8")

# === 5. Дисклеймер на слайде симуляции =======================================
итог("дисклеймер «Модельный прогноз, не замер.» на слайде симуляции",
     "PASS" if "Модельный прогноз, не замер." in тексты_слайдов[3] else "FAIL")

# === 6. Тема «графит» меняет фон контент-слайдов =============================
r2 = client.post("/api/preza/pptx",
                 data=json.dumps({"проект_id": pid, "тема": "графит"}),
                 content_type="application/json")
г_путь = OUT_DIR / "Защита_тест_графит.pptx"
г_путь.write_bytes(r2.data)
prs_g = Presentation(str(г_путь))
фон_слайда2 = prs_g.slides[1].shapes[0].fill.fore_color.rgb
фон_слайда1 = prs_g.slides[0].shapes[0].fill.fore_color.rgb
итог("тема «графит»: фон контент-слайда #2B2620 (титул тёмный в обеих темах)",
     "PASS" if str(фон_слайда2) == "2B2620" and str(фон_слайда1) == "2B2620"
     else "FAIL", f"слайд2 {фон_слайда2}, слайд1 {фон_слайда1}")
prs_b = Presentation(str(PPTX_ПУТЬ))
итог("тема «бумага»: фон контент-слайда #F7F4EF, титул тёмный",
     "PASS" if str(prs_b.slides[1].shapes[0].fill.fore_color.rgb) == "F7F4EF"
     and str(prs_b.slides[0].shapes[0].fill.fore_color.rgb) == "2B2620"
     else "FAIL")

# === 7. Чек-лист слайдов =====================================================
r3 = client.post("/api/preza/pptx",
                 data=json.dumps({"проект_id": pid,
                                  "слайды": ["титул", "эффект"]}),
                 content_type="application/json")
n3 = len(Presentation(__import__("io").BytesIO(r3.data)).slides._sldIdLst)
итог("чек-лист: запрошены 2 слайда → в файле ровно 2",
     "PASS" if r3.status_code == 200 and n3 == 2 else "FAIL",
     f"HTTP {r3.status_code}, слайдов {n3}")

# === 8. Визуальная приёмка: рендер PNG через soffice =========================
soffice = shutil.which("soffice")
if soffice:
    with tempfile.TemporaryDirectory() as td:
        subprocess.run([soffice, "--headless", "--convert-to", "pdf",
                        "--outdir", td, str(PPTX_ПУТЬ)],
                       check=True, capture_output=True, timeout=300)
        import pypdfium2 as pdfium
        pdf = pdfium.PdfDocument(str(Path(td) / (PPTX_ПУТЬ.stem + ".pdf")))
        pngs = []
        for i in range(len(pdf)):
            im = pdf[i].render(scale=1.4).to_pil()
            p = OUT_DIR / f"слайд_{i + 1:02d}.png"
            im.save(p)
            pngs.append(p)
    плохие = [p.name for p in pngs if p.stat().st_size < 5000]
    итог("рендер: 8 PNG в ui/_qa_p3_slides/, все ненулевые",
         "PASS" if len(pngs) == 8 and not плохие else "FAIL",
         f"рендеров {len(pngs)}, подозрительно малы: {плохие or '—'}")
else:
    итог("рендер: soffice не найден — визуальная приёмка пропущена", "FAIL",
         "установите LibreOffice")

# --- Свод ---
фейлы = [x for x in результаты if x[1] == "FAIL"]
print(f"\nИтого: {len(результаты) - len(фейлы)}/{len(результаты)} PASS")
sys.exit(1 if фейлы else 0)
