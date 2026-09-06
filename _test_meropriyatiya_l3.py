#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Регрессионный тест Ломтика 3 генератора мероприятий (Спека № 3, раздел 2,
шаг 6): план мероприятий по шаблону Прил. 4 + экспорт .docx/.xlsx +
контракт прогона (итоги_мероприятия / отчёт_meropriyatiya_md).

Эталон: ОКЗ — УМ «фасовка», 38,7 % ВПП (тот же вход, что в тесте ломтика 2).

Проверки:
  * plan заполнен (≥ 3 записей), обязательные колонки на месте, сортировка
    по prioritet, даты согласованы (начало ≤ окончания);
  * build_plan_docx / build_plan_xlsx — валидные байты: открываются
    python-docx / openpyxl, ключевые строки (заголовок шаблона, ИТОГО,
    сноска) на месте;
  * итоги_мероприятия отдаёт сводку, отчёт_meropriyatiya_md — markdown;
  * qc pass.

Образцы:
  ИИ-направление/результаты/план_мероприятий_v01/план_мероприятий_окз_v01.docx
  ИИ-направление/результаты/план_мероприятий_v01/план_мероприятий_окз_v01.xlsx
"""

import datetime as dt
import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from docx import Document
from openpyxl import load_workbook

from meropriyatiya import (
    сгенерировать, итоги_мероприятия, отчёт_meropriyatiya_md,
    build_plan_docx, build_plan_xlsx, _КОЛОНКИ_ПЛАНА,
)
from _test_meropriyatiya_l2 import ЭТАЛОН_ОКЗ

РЕЗУЛЬТАТЫ = (Path(__file__).resolve().parent.parent / "результаты"
              / "план_мероприятий_v01")


def main():
    r = сгенерировать(ЭТАЛОН_ОКЗ)
    ошибки = []
    план = r["plan"]

    # 1. План заполнен, обязательные колонки, сортировка по prioritet
    if len(план) < 3:
        ошибки.append(f"записей плана {len(план)} < 3")
    for з in план:
        for кол in _КОЛОНКИ_ПЛАНА:
            if з.get(кол) is None:
                ошибки.append(f"запись № {з.get('n')}: колонка «{кол}» = None")
        нач = dt.date.fromisoformat(з["data_nachala"])
        кон = dt.date.fromisoformat(з["data_okonchaniya"])
        if нач > кон:
            ошибки.append(f"запись № {з['n']}: дата начала позже окончания")
    ранги = [з["prioritet"] for з in план]
    if ранги != sorted(ранги):
        ошибки.append(f"план не отсортирован по prioritet: {ранги}")
    # № записей — сплошной ряд 1..N
    if [з["n"] for з in план] != list(range(1, len(план) + 1)):
        ошибки.append("нумерация записей плана не образует ряд 1..N")

    # 2. docx валиден
    байты_docx = build_plan_docx(ЭТАЛОН_ОКЗ, r)
    if not isinstance(байты_docx, bytes) or not байты_docx:
        ошибки.append("build_plan_docx вернул не-байты/пусто")
    else:
        try:
            doc = Document(io.BytesIO(байты_docx))
            текст = "\n".join(p.text for p in doc.paragraphs)
            if "Статус реализации плана мероприятий" not in текст:
                ошибки.append("docx: нет заголовка шаблона Прил. 4")
            if "ОКЗ" not in текст:
                ошибки.append("docx: нет реквизитов проекта (ОКЗ)")
            if "Сноска (evidence-допущения)" not in текст:
                ошибки.append("docx: нет сноски об evidence-допущениях")
            таблица = "\n".join(c.text for t in doc.tables
                                for row in t.rows for c in row.cells)
            for ключ in ("Мероприятие", "Ответственный", "Дата начала",
                         "Дата окончания", "Статус",
                         "Важно (комментарий для проверки)",
                         "ИТОГО по плану", план[0]["meropriyatie"]):
                if ключ not in таблица:
                    ошибки.append(f"docx: в таблице нет «{ключ}»")
        except Exception as e:
            ошибки.append(f"docx не открылся python-docx: {e}")

    # 3. xlsx валиден
    байты_xlsx = build_plan_xlsx(ЭТАЛОН_ОКЗ, r)
    if not isinstance(байты_xlsx, bytes) or not байты_xlsx:
        ошибки.append("build_plan_xlsx вернул не-байты/пусто")
    else:
        try:
            wb = load_workbook(io.BytesIO(байты_xlsx))
            ws = wb["План мероприятий"]
            текст = "\n".join(str(c.value) for row in ws.iter_rows()
                              for c in row if c.value is not None)
            for ключ in ("Статус реализации плана мероприятий", "ОКЗ",
                         "Мероприятие", "Ответственный", "Дата начала",
                         "Дата окончания", "Статус",
                         "Важно (комментарий для проверки)",
                         "ИТОГО по плану", "Сноска (evidence-допущения)",
                         план[0]["meropriyatie"]):
                if ключ not in текст:
                    ошибки.append(f"xlsx: нет «{ключ}»")
        except Exception as e:
            ошибки.append(f"xlsx не открылся openpyxl: {e}")

    # 4. Контракт прогона: итоги + отчёт
    итоги = итоги_мероприятия(r)
    for поле in ("мероприятий", "план_записей", "сумма_dVPP_pp",
                 "сумма_dPT_pct", "сумма_dEE_rub", "qc_pass"):
        if итоги.get(поле) is None:
            ошибки.append(f"итоги_мероприятия: поле «{поле}» = None")
    if итоги.get("план_записей") != len(план):
        ошибки.append("итоги: план_записей ≠ длине plan")
    отчёт = отчёт_meropriyatiya_md(r)
    if "## Генератор мероприятий" not in отчёт or "Топ-3 плана" not in отчёт:
        ошибки.append("отчёт_meropriyatiya_md: неполная сводка")

    # 5. QC: pass или только замечания про gap-замеры
    блокирующие = [н for н in r["qc"] if "замер" not in н and "gap" not in н]
    if блокирующие:
        ошибки.extend("qc: " + н for н in блокирующие)

    # Сохраняем образцы
    РЕЗУЛЬТАТЫ.mkdir(parents=True, exist_ok=True)
    путь_docx = РЕЗУЛЬТАТЫ / "план_мероприятий_окз_v01.docx"
    путь_xlsx = РЕЗУЛЬТАТЫ / "план_мероприятий_окз_v01.xlsx"
    путь_docx.write_bytes(байты_docx)
    путь_xlsx.write_bytes(байты_xlsx)

    print(f"Записей плана:          {len(план)}")
    print(f"Σ ΔВПП / ΔПТ / ΔЭЭ:     {итоги['сумма_dVPP_pp']} п.п. / "
          f"{итоги['сумма_dPT_pct']} % / "
          f"{итоги['сумма_dEE_rub']:,.0f} руб/год".replace(",", " "))
    print("Первые 3 записи плана:")
    for з in план[:3]:
        print(f"  {з['n']}. {з['meropriyatie']} — {з['srok']}, "
              f"класс {з['prioritet_klass']}, курс «{з['kurs_fck']}»")
    print(f"docx: {путь_docx} ({len(байты_docx)} байт)")
    print(f"xlsx: {путь_xlsx} ({len(байты_xlsx)} байт)")
    print(f"QC-замечаний:           {len(r['qc'])}")

    if ошибки:
        print("\nТЕСТ ПРОВАЛЕН:")
        for о in ошибки:
            print(" - " + о)
        return 1
    print("\nТЕСТ ПРОЙДЕН")
    return 0


if __name__ == "__main__":
    sys.exit(main())
