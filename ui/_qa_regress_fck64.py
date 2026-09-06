# -*- coding: utf-8 -*-
"""Регрессия Разработчик_ФЦК-64: все инструменты + новые akt/soglashenie."""
import importlib.util
import io
import json
import sqlite3
import urllib.request
import urllib.error
from pathlib import Path

from docx import Document

import os
BASE = os.environ.get("AGENT_CEH_BASE", "http://localhost:8000")
AUTH = ("рцк", "H8BgJlzfjESRDv4F")
UI = Path(__file__).resolve().parent
PROT = UI.parent

результаты = []


def итог(инструмент, статус, эталон=""):
    результаты.append((инструмент, статус, эталон))
    print(f"[{статус}] {инструмент} {эталон}")


def req(method, path, body=None, expect_json=True):
    import base64
    import urllib.parse
    path = urllib.parse.quote(path, safe="/?=&")
    r = urllib.request.Request(BASE + path, method=method)
    r.add_header("Authorization", "Basic " + base64.b64encode(
        f"{AUTH[0]}:{AUTH[1]}".encode()).decode())
    data = None
    if body is not None:
        data = json.dumps(body, ensure_ascii=False).encode()
        r.add_header("Content-Type", "application/json")
    try:
        with urllib.request.urlopen(r, data=data, timeout=60) as resp:
            blob = resp.read()
            ct = resp.headers.get("Content-Type", "")
            if expect_json and "json" in ct:
                return json.loads(blob.decode()), None
            return None, blob
    except urllib.error.HTTPError as e:
        return {"ok": False, "errors": [f"HTTP {e.code}: {e.read()[:300]}"]}, None


# --- UM (ОКЗ): эталон — главное УМ «фасовка», 38,7 % ВПП ---
demo, _ = req("GET", "/api/demo")
res, _ = req("POST", "/api/raschet", demo)
ит = res.get("итоги") or res.get("итоги_json") or {}
ум = (ит.get("главное_ум") or json.dumps(res, ensure_ascii=False))
ок = res.get("ok") and "фасовк" in json.dumps(res, ensure_ascii=False)
доля_ок = "38,7" in json.dumps(res, ensure_ascii=False).replace(".", ",")
итог("UM (ОКЗ)", "PASS" if ок and доля_ок else "FAIL",
     f"УМ «фасовка»: {'да' if ок else 'НЕТ'}; 38,7 % ВПП: {'да' if доля_ок else 'НЕТ'}")

# --- VSM: тот же вход + glavnoe ---
vsm_in = dict(demo); vsm_in["glavnoe"] = "фасовка"
res, _ = req("POST", "/api/vsm", vsm_in)
ок = res.get("ok") and res.get("svg") and "<svg" in res["svg"]
итог("vsm", "PASS" if ок else "FAIL",
     f"svg: {'да' if ок else json.dumps(res, ensure_ascii=False)[:120]}")

# --- ЭЭ (НОВОХРОМ): эталон 23 757 960,67 ₽ ---
ee_in, _ = req("GET", "/api/ee_demo/новохром")
res, _ = req("POST", "/api/ee_raschet", ee_in)
raw = json.dumps(res, ensure_ascii=False)
ок = res.get("ok") and "23 757 960,67" in raw
итог("ЭЭ (НОВОХРОМ)", "PASS" if ок else "FAIL",
     f"ээ совпал с эталоном 23 757 960,67: {'да' if ок else raw[:150]}")

# --- ee_export_fck: справка ФЦК (xlsm) ---
res, blob = req("POST", "/api/ee_export_fck", ee_in, expect_json=False)
ок = blob is not None and blob[:2] == b"PK"
итог("ee_export_fck", "PASS" if ок else "FAIL",
     f"xlsm: {'да, %d байт' % len(blob) if ок else str(res)[:120]}")

# --- protokol: эталонные данные Кушкуль ---
with open(PROT / "данные_протокол_кушкуль.json", encoding="utf-8") as f:
    pz = json.load(f)
res, blob = req("POST", "/api/protokol", {"данные": pz}, expect_json=False)
ок = False
if blob and blob[:2] == b"PK":
    doc = Document(io.BytesIO(blob))
    txt = "\n".join(p.text for p in doc.paragraphs)
    ок = "Кушкуль" in txt or "Протокол" in txt
итог("protokol", "PASS" if ок else "FAIL",
     f"docx с реквизитами: {'да' if ок else str(res)[:120]}")

# --- ord: вход из прогона № 14 ---
con = sqlite3.connect(UI / "agent_ceh.db")
ord_in = json.loads(con.execute(
    "SELECT вход_json FROM progony WHERE id=14").fetchone()[0])
sm_in = json.loads(con.execute(
    "SELECT вход_json FROM progony WHERE id=12").fetchone()[0])
con.close()
res, blob = req("POST", "/api/ord", ord_in, expect_json=False)
ок = False
if blob and blob[:2] == b"PK":
    doc = Document(io.BytesIO(blob))
    txt = "\n".join(p.text for p in doc.paragraphs)
    ок = "приказ" in txt.lower() or "карточка" in txt.lower()
итог("ord", "PASS" if ок else "FAIL", f"docx: {'да' if ок else str(res)[:120]}")

# --- smart: вход из прогона № 12 ---
res, _ = req("POST", "/api/smart", sm_in)
ок = res.get("ok") and res.get("итоги")
итог("smart", "PASS" if ок else "FAIL",
     f"целей: {res.get('итоги', {}).get('целей') if ок else str(res)[:120]}")

# --- otbor: РЦК 19 баллов → «проходит»; ФЦК ---
s = importlib.util.spec_from_file_location("otbor", PROT / "otbor.py")
otbor = importlib.util.module_from_spec(s); s.loader.exec_module(otbor)
форм = {k: True for k, _ in otbor.ФОРМАЛЬНЫЕ_КРИТЕРИИ}

def вход_отбор(трек, да):
    сп = otbor.СПЕКИ_ТРЕКОВ[трек]
    пункты = list(сп["пункты"].keys())
    отсек = [k for k, v in сп["пункты"].items() if v.get("отсекающий")]
    ответы = []
    n = 0
    for k in пункты:
        val = k in отсек or n < да - len([x for x in отсек if x in пункты])
        if k in отсек:
            val = True
        else:
            val = n < да - len(отсек)
            if val:
                n += 1
        ответы.append({"пункт": k, "ответ": bool(val)})
    return {"предприятие": "ООО «Регрессия»", "трек": трек,
            "формальная_проверка": форм, "ответы_чеклиста": ответы}

res, _ = req("POST", "/api/otbor", вход_отбор("рцк", 19))
r = res.get("результат", {})
баллы = r.get("чеклист", {}).get("баллы")
вердикт = r.get("вердикт", "")
ок = res.get("ok") and баллы == 19 and "проход" in вердикт.lower()
итог("otbor (РЦК)", "PASS" if ок else "FAIL",
     f"баллы={баллы}, вердикт: {вердикт[:60]}")

res, _ = req("POST", "/api/otbor", вход_отбор("фцк", 20))
r = res.get("результат", {})
баллы = r.get("чеклист", {}).get("баллы")
ок = res.get("ok") and баллы is not None
итог("otbor (ФЦК)", "PASS" if ок else "FAIL",
     f"баллы={баллы}, вердикт: {r.get('вердикт', '')[:60]}")

# --- otbor_potok: 2 кандидата ---
# берём критерии из ядра
s2 = importlib.util.spec_from_file_location("otbor_potok", PROT / "otbor_potok.py")
op = importlib.util.module_from_spec(s2); s2.loader.exec_module(op)
критерии = [k for k, _, _ in op.КРИТЕРИИ]
признаки = [k for k, _ in op.ПРИЗНАКИ_ПОТЕРЬ]
potok_in = {"предприятие": "ООО «Регрессия»", "кандидаты": [
    {"название": "поток А",
     "критерии": {k: True for k in критерии},
     "признаки": {k: True for k in признаки[:10]}},
    {"название": "поток Б",
     "критерии": {k: (i % 2 == 0) for i, k in enumerate(критерии)},
     "признаки": {k: True for k in признаки[:4]}},
]}
res, _ = req("POST", "/api/otbor/potok", potok_in)
r = res.get("результат", {})
ок = res.get("ok") and r.get("рекомендованный") == "поток А"
итог("potok", "PASS" if ок else "FAIL",
     f"рекомендованный: {r.get('рекомендованный') or str(res)[:120]}")

# --- sravnenie: прогоны УМ 9 и 11 ---
res, _ = req("POST", "/api/sravnenie", {"id_do": 9, "id_posle": 11})
ок = res.get("ok")
итог("sravnenie", "PASS" if ок else "FAIL",
     ("дельта получена" if ок else str(res)[:120]))

# --- akt: генерация docx + валидность ---
akt_in = {"вариант": "рцк",
          "центр": "АНО «Центр поддержки предпринимательства Оренбургской области»",
          "предприятие": "ООО «Регрессия»", "город": "г. Оренбург",
          "дата_акта": "«05» августа 2026 г.",
          "соглашение": {"номер": "12", "дата": "«12» октября 2025 г."},
          "дата_начала": "«01» сентября 2026 г.",
          "ответственный_центра": "Судакова О.С., руководитель проектного офиса",
          "подписанты": {"предприятие": {"должность": "Директор", "фио": "Иванов И.И."},
                         "центр": {"должность": "Руководитель РЦК", "фио": "Судакова О.С."}}}
res, blob = req("POST", "/api/akt", akt_in, expect_json=False)
ок = False
if blob and blob[:2] == b"PK":
    doc = Document(io.BytesIO(blob))
    txt = "\n".join(p.text for p in doc.paragraphs)
    табл = "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    весь = txt + табл
    ок = "Регрессия" in весь and "01" in весь and "сентября" in весь
итог("akt (новый)", "PASS" if ок else "FAIL",
     f"docx, реквизиты: {'да' if ок else str(res)[:150]}")

# --- soglashenie ---
sg_in = {"номер": "12", "место": "г. Оренбург", "дата": "«12» октября 2025 г.",
         "предприятие": "ООО «Регрессия»",
         "представитель_предприятия": "генерального директора Иванова Ивана Ивановича",
         "рцк": "АНО «Центр поддержки предпринимательства Оренбургской области»",
         "представитель_рцк": "руководителя Судаковой Ольги Сергеевны",
         "основание_рцк": "устава",
         "штраф_отчетность": "10 000 (десять тысяч)",
         "штраф_расторжение": "100 000 (сто тысяч)",
         "реквизиты": {"предприятие": {"наименование": "ООО «Регрессия»",
                                       "инн_кпп": "7700000000 / 770001001"}},
         "подписанты": {"предприятие": {"фио": "Иванов И.И."},
                        "рцк": {"фио": "Судакова О.С."}}}
res, blob = req("POST", "/api/soglashenie", sg_in, expect_json=False)
ок = False
if blob and blob[:2] == b"PK":
    doc = Document(io.BytesIO(blob))
    весь = "\n".join(p.text for p in doc.paragraphs) + \
        "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    ок = "Соглашение о сотрудничестве" in весь and "Регрессия" in весь \
        and "Иванов" in весь
итог("soglashenie (новый)", "PASS" if ок else "FAIL",
     f"docx, реквизиты: {'да' if ок else str(res)[:150]}")

# --- сохранение прогонов akt/soglashenie в проект 2 (контракт) ---
res, _ = req("POST", "/api/progony", {"id_proekta": 2, "инструмент": "akt",
             "вход": akt_in, "автор": "РЦК", "комментарий": "регрессия ФЦК-64: акт"})
ок1 = res.get("ok") and res.get("итоги", {}).get("файл")
gid_akt = res.get("id")
res, _ = req("POST", "/api/progony", {"id_proekta": 2, "инструмент": "soglashenie",
             "вход": sg_in, "автор": "РЦК", "комментарий": "регрессия ФЦК-64: соглашение"})
ок2 = res.get("ok") and res.get("итоги", {}).get("файл")
gid_sg = res.get("id")
итог("сохранение прогонов akt/soglashenie",
     "PASS" if ок1 and ок2 else "FAIL",
     f"akt № {gid_akt} файл={bool(ок1)}, соглашение № {gid_sg} файл={bool(ок2)}")

# --- скачивание файла прогона ---
if gid_akt:
    res, blob = req("GET", f"/api/progony/{gid_akt}/файл", expect_json=False)
    ок = blob is not None and blob[:2] == b"PK"
    итог("скачивание файла прогона akt", "PASS" if ок else "FAIL",
         f"progon_{gid_akt}_акт.docx: {'да' if ок else str(res)[:100]}")

print()
плохие = [r for r in результаты if r[1] != "PASS"]
print(f"ИТОГО: {len(результаты) - len(плохие)}/{len(результаты)} PASS")
for r in плохие:
    print("ПРОВАЛ:", r)
