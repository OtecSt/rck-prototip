# -*- coding: utf-8 -*-
"""PPTX-движок фаст-презентации (Спека П1, ломтик П3).

Превращает выход сборщика П2 (preza_data.собрать) в PPTX 16:9.
Приёмы вёрстки переняты из референса вкр_находки/build_premium.py
(slide/box/rect/stat/kicker/title/shp_text, сетка 0,7", шаг карточек 3,05–4,1),
палитра — DESIGN_V2 (темы «бумага»/«графит»), данные ВКР/100У не используются.
Без morph-переходов, без внешних картинок, без LLM: вся графика — фигуры
python-pptx из чисел прогонов. Каждое число форматируется по-русски
(пробелы тысяч, запятая десятичная) — только здесь, в П2 значения сырые.
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

try:
    from pptx.enum.line import MSO_LINE_DASH_STYLE
    _ПУНКТИР = MSO_LINE_DASH_STYLE.DASH
except Exception:  # старые версии python-pptx — пунктир недоступен
    _ПУНКТИР = None

# --- Палитра DESIGN_V2 (Спека П1, раздел 5) ---
БУМАГА = {
    "фон": RGBColor(0xF7, 0xF4, 0xEF),
    "карточка": RGBColor(0xFF, 0xFF, 0xFF),
    "бордер": RGBColor(0xE8, 0xE2, 0xD6),
    "текст": RGBColor(0x2B, 0x26, 0x20),
    "вторичный": RGBColor(0x6E, 0x67, 0x5B),
    "приглушённый": RGBColor(0x9A, 0x91, 0x84),
    "акцент": RGBColor(0xB4, 0x56, 0x2F),
    "акцент_фон": RGBColor(0xF5, 0xE9, 0xE0),
    "успех": RGBColor(0x4E, 0x7A, 0x51),
    "потери": RGBColor(0xA8, 0x43, 0x3A),
}
ГРАФИТ = dict(БУМАГА)
ГРАФИТ.update({
    "фон": RGBColor(0x2B, 0x26, 0x20),
    "карточка": RGBColor(0x3A, 0x34, 0x2C),
    "бордер": RGBColor(0x4A, 0x43, 0x39),
    "текст": RGBColor(0xF7, 0xF4, 0xEF),
    "вторичный": RGBColor(0xC7, 0xBF, 0xB2),
    "приглушённый": RGBColor(0x9A, 0x91, 0x84),
    "акцент": RGBColor(0xC9, 0x6A, 0x43),
    "акцент_фон": RGBColor(0x45, 0x36, 0x2C),
})

ТЕМЫ = {"бумага": БУМАГА, "графит": ГРАФИТ}
F = "Calibri"  # единственный шрифт (Спека П1, п. 5: гарантированно везде)

ШИРИНА, ВЫСОТА = 13.333, 7.5
ПОЛЕ = 0.7


# --- Форматирование чисел по-русски (единственное место форматирования) ---
def fmt_деньги(x):
    """23 757 960,67"""
    if x is None:
        return "—"
    return f"{x:,.2f}".replace(",", " ").replace(".", ",")


def fmt_число(x, nd=2):
    """2,00 · 0,704167 (nd по правилам инструмента-источника)"""
    if x is None:
        return "—"
    if isinstance(x, float) and x.is_integer() and nd == 0:
        return f"{int(x):,}".replace(",", " ")
    s = f"{x:,.{nd}f}".replace(",", " ").replace(".", ",")
    return s


def fmt_пкт(x, nd=1):
    return "—" if x is None else fmt_число(x, nd) + " %"


def fmt_дельта(x, nd=0):
    if x is None:
        return "—"
    знак = "+" if x >= 0 else "−"
    return знак + fmt_число(abs(x), nd)


# --- Базовые приёмы вёрстки (мотивы build_premium.py) ---
class _Лист:
    """Один слайд + активная тема: обёртка над приёмами, чтобы не таскать П."""

    def __init__(self, s, п):
        self.s = s
        self.п = п

    def box(self, x, y, w, h, t, sz, цвет=None, bold=False, al=PP_ALIGN.LEFT,
            it=False, anch=MSO_ANCHOR.TOP):
        цвет = цвет or self.п["текст"]
        tb = self.s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anch
        for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, m, 0)
        for i, ln in enumerate(str(t).split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = al
            r = p.add_run()
            r.text = ln
            r.font.size = Pt(sz)
            r.font.bold = bold
            r.font.italic = it
            r.font.color.rgb = цвет
            r.font.name = F
        return tb

    def rect(self, x, y, w, h, fill, line=None, rounded=True):
        shp = self.s.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
            Inches(x), Inches(y), Inches(w), Inches(h))
        if rounded:
            try:
                shp.adjustments[0] = 0.08
            except Exception:
                pass
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
        if line:
            shp.line.color.rgb = line
            shp.line.width = Pt(1.0)
        else:
            shp.line.fill.background()
        shp.shadow.inherit = False
        return shp

    def shp_text(self, shp, lines, al=PP_ALIGN.CENTER):
        tf = shp.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for i, (txt, sz, цвет, bold) in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = al
            r = p.add_run()
            r.text = txt
            r.font.size = Pt(sz)
            r.font.color.rgb = цвет
            r.font.bold = bold
            r.font.name = F

    def пунктир(self, x1, y1, x2, y2, цвет):
        """Тонкая пунктирная линия (линия такта)."""
        from pptx.enum.shapes import MSO_CONNECTOR
        ln = self.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                         Inches(x1), Inches(y1),
                                         Inches(x2), Inches(y2))
        ln.line.color.rgb = цвет
        ln.line.width = Pt(1.0)
        if _ПУНКТИР is not None:
            ln.line.dash_style = _ПУНКТИР
        ln.shadow.inherit = False
        return ln

    def kicker(self, t):
        self.box(ПОЛЕ, 0.5, 12, 0.4, t.upper(), 12, self.п["вторичный"], bold=True)

    def title(self, t, sz=29):
        self.box(ПОЛЕ, 0.95, 11.93, 1.0, t, sz, self.п["текст"], bold=True)

    def stat(self, x, y, w, num, lab, цвет_числа=None, h=1.5, sz=26, sz_lab=12):
        цвет_числа = цвет_числа or self.п["текст"]
        self.rect(x, y, w, h, self.п["карточка"], line=self.п["бордер"])
        self.box(x + 0.12, y + 0.22, w - 0.24, 0.75, num, sz, цвет_числа,
                 bold=True, al=PP_ALIGN.CENTER)
        self.box(x + 0.12, y + h - 0.48, w - 0.24, 0.42, lab, sz_lab,
                 self.п["вторичный"], al=PP_ALIGN.CENTER)

    def сноска(self, t):
        self.box(ПОЛЕ, 7.02, 11.93, 0.35, t, 10.5, self.п["приглушённый"], it=True)


def _новый_слайд(prs, п, тёмный=False):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    пал = ГРАФИТ if тёмный else п
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width,
                           prs.slide_height)
    r.fill.solid()
    r.fill.fore_color.rgb = пал["фон"]
    r.line.fill.background()
    r.shadow.inherit = False
    return _Лист(s, пал)


def _заметки(s, слайд):
    """Speaker notes: трассировка чисел к прогонам (стоп-гейт спеки, п. 7)."""
    if слайд["источники"]:
        строки = [f"источник: прогон № {и['прогон_id']} "
                  f"({и['инструмент']}, {и['дата']})"
                  for и in слайд["источники"]]
    else:
        строки = ["источник: карточка проекта"]
    s.notes_slide.notes_text_frame.text = "\n".join(строки)


def _полоса_метрик(л, y, метрики, шаг=3.05, w=2.9):
    """метрики: [(подпись, значение, цвет_числа или None)] — максимум 4 в ряд."""
    x = ПОЛЕ
    for lab, val, цвет in метрики[:4]:
        л.stat(x, y, w, val, lab, цвет_числа=цвет)
        x += шаг


# --- Шаблоны слайдов ---
def _слайд_титул(prs, п, слайд, прогоны_данные):
    л = _новый_слайд(prs, п, тёмный=True)
    d = слайд["данные"]
    л.box(0.9, 1.55, 11.5, 0.5, "ПИЛОТНЫЙ ПРОЕКТ · РЦК", 13,
          ГРАФИТ["вторичный"], bold=True)
    л.box(0.9, 2.1, 11.5, 1.6, d.get("предприятие") or d.get("название_проекта")
          or "Пилотный проект", 38, ГРАФИТ["текст"], bold=True)
    if d.get("поток"):
        л.box(0.9, 3.55, 11.5, 0.6, f"Поток: {d['поток']}", 17, ГРАФИТ["акцент"],
              bold=True)
    строка = " · ".join(x for x in (d.get("автор"), d.get("дата")) if x)
    if строка:
        л.box(0.9, 4.3, 11.5, 0.5, строка, 14, ГРАФИТ["вторичный"])

    # Фирменная обложка «каналы потока»: нижняя треть — столбцы, высоты ∝
    # ряду операций из прогона диагностики (доли_впп: имя, сек, доля);
    # красный — главное узкое место. Нет ряда — равные столбцы.
    ряд, ум = [], None
    проб = прогоны_данные.get("проблема")
    if проб:
        ряд = [(имя, сек) for имя, сек, _ in (проб["данные"].get("доли_впп") or [])]
        ум = проб["данные"].get("главное_ум")
    n = max(len(ряд), 1)
    if not ряд:
        ряд = [("", 1.0)] * 8
        n = 8
    макс = max(v for _, v in ряд) or 1.0
    x0, ширина_полосы, gap = 0.9, 11.53, 0.14
    bw = (ширина_полосы - (n - 1) * gap) / n
    y_base, h_max = 6.9, 1.5
    for i, (имя, v) in enumerate(ряд):
        h = max(0.12, h_max * v / макс)
        цвет = БУМАГА["потери"] if имя and имя == ум else ГРАФИТ["акцент"]
        л.rect(x0 + i * (bw + gap), y_base - h, bw, h, цвет, rounded=False)
    л.box(0.9, 7.0, 11.5, 0.35,
          "каналы потока: высота столбца — время операции в потоке, "
          "красный — главное узкое место" if ум else
          "каналы потока проекта", 10, ГРАФИТ["приглушённый"], it=True)
    return л


def _слайд_проблема(prs, п, слайд):
    л = _новый_слайд(prs, п)
    d = слайд["данные"]
    л.kicker(слайд["кикер"])
    л.title(слайд["заголовок"])
    ум = d.get("главное_ум") or "—"
    имя_ум = ум if len(ум) <= 58 else ум[:55] + "…"
    метрики = [
        ("Главное узкое место", имя_ум, п["потери"]),
        ("Доля во ВПП", fmt_пкт(d.get("доля_впп_ум_пкт"), 2), п["потери"]),
        ("Статус", ("подтверждено" if d.get("подтверждено")
                    else "не подтверждено"),
         п["успех"] if d.get("подтверждено") else п["вторичный"]),
        ("Операций в разборе", str(len(d.get("доли_впп") or [])), None),
    ]
    # Первая карточка с именем УМ — текстом меньше, чтобы не лезла за край
    x = ПОЛЕ
    for i, (lab, val, цвет) in enumerate(метрики):
        sz = 15 if i == 0 else (14 if i == 2 else 26)
        л.stat(x, 2.15, 2.9, val, lab, цвет_числа=цвет, sz=sz)
        x += 3.05
    ранж = d.get("ранжирование") or []
    if ранж:
        л.box(ПОЛЕ, 4.05, 11.93, 0.4, "Ранжирование узких мест", 14,
              п["вторичный"], bold=True)
        строки = "\n".join(f"{i + 1}.  {имя}" for i, имя in enumerate(ранж[:5]))
        л.box(ПОЛЕ, 4.5, 11.93, 2.3, строки, 15, п["текст"])
    return л


def _слайд_поток(prs, п, слайд):
    л = _новый_слайд(prs, п)
    d = слайд["данные"]
    п_ = d["поток"]
    л.kicker(слайд["кикер"])
    л.title(слайд["заголовок"])
    _полоса_метрик(л, 2.1, [
        ("Такт, мин/шт", fmt_число(п_.get("такт_мин")), None),
        ("ВПП, мин", fmt_число(п_.get("впп_мин")), п["акцент"]),
        ("Время цикла ΣТц, мин", fmt_число(п_.get("сум_тц_мин")), None),
        ("Выработка, шт/чел", fmt_число(п_.get("выработка_шт_на_чел")), None),
    ])
    строки = []
    if п_.get("узкое_место"):
        строки.append(f"Узкое место по расчёту: «{п_['узкое_место']}» "
                      f"(Тц {fmt_число(п_.get('узкое_место_тц_мин'))} мин)")
    выше = п_.get("выше_такта") or []
    if выше:
        строки.append("Выше такта: " + ", ".join(f"«{x}»" for x in выше))
    if п_.get("потенциал_впп_мин") is not None:
        строки.append("Потенциал сокращения ВПП: до "
                      f"{fmt_число(п_['потенциал_впп_мин'])} мин")
    if п_.get("операторов_по_норме") is not None:
        строки.append(f"Операторов по норме: {fmt_число(п_['операторов_по_норме'])}")
    if строки:
        л.box(ПОЛЕ, 4.0, 11.93, 1.5, "\n".join("•  " + s for s in строки),
              14.5, п["текст"])
    oee = d.get("oee")
    if oee:
        л.box(ПОЛЕ, 5.35, 11.93, 0.4, "OEE оборудования", 14,
              п["вторичный"], bold=True)
        _полоса_метрик(л, 5.75, [
            ("OEE", fmt_число(oee.get("OEE"), 6), п["акцент"]),
            ("Доступность A", fmt_число(oee.get("A"), 6), None),
            ("Производительность P", fmt_число(oee.get("P"), 6), None),
            ("Качество Q", fmt_число(oee.get("Q"), 6), None),
        ], шаг=3.05, w=2.9)
    return л


def _слайд_симуляция(prs, п, слайд):
    л = _новый_слайд(prs, п)
    d = слайд["данные"]
    л.kicker(слайд["кикер"])
    л.title(слайд["заголовок"])
    дельта = d.get("дельта")
    if дельта and дельта.get("числа"):
        ум = дельта.get("узкое_место") or {}
        if ум.get("переехало"):
            баннер = л.rect(ПОЛЕ, 1.95, 11.93, 0.62, п["акцент_фон"],
                            line=п["акцент"])
            л.shp_text(баннер, [(f"Узкое место переехало: «{ум.get('до')}» → "
                                 f"«{ум.get('после')}»", 14, п["акцент"], True)])
        y0 = 2.8 if ум.get("переехало") else 2.3
        показатели = [("выпуск_шт", 0), ("впп_мин", 2), ("нзп_среднее_шт", 1)]
        ряды = [(k, дельта["числа"][k], nd) for k, nd in показатели
                if k in дельта["числа"]]
        y = y0
        for k, v, nd in ряды:
            макс = max(abs(v["до"]), abs(v["после"])) or 1.0  # шкала по строке
            л.box(ПОЛЕ, y, 3.3, 0.4, v["подпись"], 13.5, п["вторичный"],
                  bold=True)
            # дельта-бары «база → сценарий»: пара горизонтальных полос
            полоса_x, полоса_w = 4.3, 5.6
            h_бар = 0.24
            л.rect(полоса_x, y + 0.02, max(0.1, полоса_w * abs(v["до"]) / макс),
                   h_бар, п["приглушённый"], rounded=False)
            л.rect(полоса_x, y + 0.3, max(0.1, полоса_w * abs(v["после"]) / макс),
                   h_бар, п["акцент"], rounded=False)
            л.box(10.15, y + 0.02, 2.5, 0.55,
                  f"{fmt_число(v['до'], nd)} → {fmt_число(v['после'], nd)}  "
                  f"({fmt_дельта(v['дельта'], nd)})", 13, п["текст"], bold=True)
            y += 0.78
        л.box(4.3, y + 0.05, 8, 0.35,
              "серая полоса — «до», терракота — «после»", 11,
              п["приглушённый"], it=True)
    else:
        б = d.get("база") or {}
        _полоса_метрик(л, 2.3, [
            ("Выпуск за горизонт, шт", fmt_число(б.get("выпуск_шт"), 0),
             п["акцент"]),
            ("ВПП, мин", fmt_число(б.get("впп_мин")), None),
            ("НЗП среднее, шт", fmt_число(б.get("нзп_среднее_шт"), 1), None),
            ("Такт, мин/шт", fmt_число(б.get("такт_мин")), None),
        ])
        if б.get("узкое_место"):
            л.box(ПОЛЕ, 4.2, 11.93, 0.5,
                  f"Узкое место в модели: «{б['узкое_место']}»", 15, п["текст"])
    л.сноска(d.get("дисклеймер") or "Модельный прогноз, не замер.")
    return л


def _топ_плана(отчёт_md):
    """Строки «Топ-3 плана» из отчёта прогона мероприятий (как хранит снапшот)."""
    строки, взять = [], False
    for ln in (отчёт_md or "").splitlines():
        if ln.startswith("### Топ-3 плана"):
            взять = True
            continue
        if взять:
            if ln.strip().startswith("#"):
                break
            if ln.strip() and ln.lstrip()[0].isdigit():
                строки.append(ln.strip())
    return строки


def _слайд_мероприятия(prs, п, слайд):
    л = _новый_слайд(prs, п)
    d = слайд["данные"]
    л.kicker(слайд["кикер"])
    л.title(слайд["заголовок"])
    _полоса_метрик(л, 2.05, [
        ("Мероприятий", str(d.get("мероприятий") or 0), п["акцент"]),
        ("Класс A / B / C",
         f"{d.get('класс_A', 0)} / {d.get('класс_B', 0)} / {d.get('класс_C', 0)}",
         None),
        ("Σ ΔВПП, п.п.", fmt_число(d.get("сумма_dVPP_pp")), None),
        ("Σ ΔЭЭ, руб/год", fmt_деньги(d.get("сумма_dEE_rub")), п["успех"]),
    ])
    топ = _топ_плана(d.get("отчёт_md"))
    if топ:
        л.box(ПОЛЕ, 3.95, 11.93, 0.4, "Приоритетные мероприятия", 14,
              п["вторичный"], bold=True)
        л.box(ПОЛЕ, 4.4, 11.93, 2.2, "\n".join(топ[:5]), 13.5, п["текст"])
    if d.get("участок"):
        л.сноска(f"Узкое место: участок «{d['участок']}»"
                 + (f", ВПП {fmt_пкт(d.get('vpp_pct'))}"
                    if d.get("vpp_pct") is not None else ""))
    return л


def _слайд_эффект(prs, п, слайд):
    л = _новый_слайд(prs, п)
    d = слайд["данные"]
    л.kicker(слайд["кикер"])
    л.title(слайд["заголовок"])
    слои = [("Реальный ЭЭ, руб/год", d.get("ээ"), п["акцент"]),
            ("Потенциальный ЭЭ, руб/год", d.get("ээ_пот"), None),
            ("Высвобождение ДС, руб", d.get("высвоб"), None),
            ("Налоги за 3 года, руб", d.get("налоги"), None)]
    слои = [(lab, v, ц) for lab, v, ц in слои if v is not None]
    # Слои эффекта ступенями: высота карточки ∝ величине
    if слои:
        макс = max(abs(v) for _, v, _ in слои) or 1.0
        n = len(слои)
        шаг = 12.0 / n
        w = шаг - 0.15
        x = ПОЛЕ
        y_base = 5.3
        for lab, v, ц in слои:
            h = 0.9 + 2.1 * abs(v) / макс
            л.rect(x, y_base - h, w, h, п["карточка"], line=п["бордер"])
            л.box(x + 0.15, y_base - h + 0.18, w - 0.3, 0.5, lab, 12,
                  п["вторичный"])
            л.box(x + 0.15, y_base - h + 0.68, w - 0.3, 0.75,
                  fmt_деньги(v), 21, ц or п["текст"], bold=True)
            x += шаг
    дельта = d.get("дельта")
    if дельта and (дельта.get("числа") or {}).get("ээ"):
        v = дельта["числа"]["ээ"]
        л.box(ПОЛЕ, 5.55, 11.93, 0.5,
              f"Пересчёт «до/после»: {fmt_деньги(v['до'])} → "
              f"{fmt_деньги(v['после'])} руб/год "
              f"(Δ {fmt_дельта(v['дельта'], 2)})", 14, п["текст"], bold=True)
    статусы = d.get("статусы_допущения")
    if статусы:
        строки = []
        for k, v in статусы.items():
            if k == "контекст_расходится" and v:
                строки.append("•  контекст прогона расходится с карточкой "
                              "проекта — цифры относятся к входу прогона")
            elif isinstance(v, list):
                строки.extend(f"•  {x}" for x in v)
            elif isinstance(v, bool):
                строки.append(f"•  {k}: {'да' if v else 'нет'}")
            else:
                строки.append(f"•  {k}: {v}")
        л.box(ПОЛЕ, 6.1, 11.93, 0.8, "\n".join(строки[:4]), 11,
              п["приглушённый"], it=True)
    л.сноска("Суммы — из прогона ЭЭ-калькулятора, без пересчёта.")
    return л


def _слайд_план(prs, п, слайд):
    л = _новый_слайд(prs, п)
    d = слайд["данные"]
    л.kicker(слайд["кикер"])
    л.title(слайд["заголовок"])
    y = 2.0
    smart = d.get("smart")
    if smart and smart.get("цели"):
        л.box(ПОЛЕ, y, 7.2, 0.4, "Цели SMART", 14, п["вторичный"], bold=True)
        y += 0.45
        for ц in smart["цели"][:6]:
            части = [ц.get("наименование") or "цель"]
            if ц.get("текущий") is not None and ц.get("целевой") is not None:
                части.append(f"{fmt_число(ц['текущий'])} → "
                             f"{fmt_число(ц['целевой'])} {ц.get('ед') or ''}"
                             .strip())
            if ц.get("срок"):
                части.append(f"срок: {ц['срок']}")
            л.box(ПОЛЕ, y, 7.2, 0.55, "•  " + " · ".join(части), 13.5,
                  п["текст"])
            y += 0.58
    протокол = d.get("protokol")
    if протокол:
        x = 8.35
        л.box(x, 2.0, 4.28, 0.4, "Статус по протоколу закрытия", 14,
              п["вторичный"], bold=True)
        баллы = протокол.get("итого_баллов")
        порог = протокол.get("порог")
        if баллы is not None and порог is not None:
            л.stat(x, 2.45, 4.28, f"{баллы} / {порог}", "баллов из порога",
                   цвет_числа=(п["успех"] if протокол.get("порог_пройден")
                               else п["потери"]))
            статус = ("порог пройден — целевой уровень достигнут"
                      if протокол.get("порог_пройден")
                      else "порог не пройден — работа продолжается")
            л.box(x, 4.1, 4.28, 0.8, статус, 13, п["текст"])
        if протокол.get("ээ_руб_год") is not None:
            л.box(x, 5.0, 4.28, 0.6,
                  f"ЭЭ по протоколу: {fmt_деньги(протокол['ээ_руб_год'])} руб/год",
                  13, п["акцент"], bold=True)
    return л


def _слайд_выводы(prs, п, слайд):
    л = _новый_слайд(prs, п, тёмный=True)
    d = слайд["данные"]
    л.box(0.9, 0.8, 11.5, 0.9, слайд["заголовок"], 32, ГРАФИТ["текст"],
          bold=True)
    буллеты = d.get("буллеты") or []
    л.box(0.9, 1.9, 11.5, 4.3, "\n\n".join("•  " + б for б in буллеты), 16,
          ГРАФИТ["текст"])
    л.box(0.9, 6.5, 11.5, 0.5, d.get("финал") or "Спасибо за внимание", 18,
          ГРАФИТ["акцент"], bold=True, al=PP_ALIGN.CENTER)
    return л


_ШАБЛОНЫ = {
    "проблема": _слайд_проблема,
    "поток": _слайд_поток,
    "симуляция": _слайд_симуляция,
    "мероприятия": _слайд_мероприятия,
    "эффект": _слайд_эффект,
    "план": _слайд_план,
}


def собрать_pptx(собрано, путь):
    """Собрать PPTX 16:9 из выхода сборщика П2 и сохранить в путь."""
    тема = собрано.get("тема") or "бумага"
    п = ТЕМЫ.get(тема, БУМАГА)
    prs = Presentation()
    prs.slide_width = Inches(ШИРИНА)
    prs.slide_height = Inches(ВЫСОТА)
    прогоны_данные = {s["код"]: s for s in собрано["слайды"]}
    for слайд in собрано["слайды"]:
        код = слайд["код"]
        if код == "титул":
            л = _слайд_титул(prs, п, слайд, прогоны_данные)
        elif код == "выводы":
            л = _слайд_выводы(prs, п, слайд)
        else:
            л = _ШАБЛОНЫ[код](prs, п, слайд)
        _заметки(л.s, слайд)
    prs.save(str(путь))
    return Path(путь)


def проверить_pptx(путь, ожидали_слайдов):
    """Round-trip валидация (Спека П1, п. 7): файл открывается python-pptx,
    размер ровно 16:9, слайдов столько, сколько собрано, шрифт Calibri.
    Возвращает список ошибок (пустой = ок)."""
    ошибки = []
    try:
        prs = Presentation(str(путь))
    except Exception as e:
        return [f"собранный PPTX не открывается python-pptx: {e}"]
    w, h = prs.slide_width, prs.slide_height
    if (w, h) != (Inches(ШИРИНА), Inches(ВЫСОТА)):
        ошибки.append(f"размер слайдов {w}×{h} EMU — не 16:9 "
                      f"({Inches(ШИРИНА)}×{Inches(ВЫСОТА)})")
    n = len(prs.slides._sldIdLst)
    if n != ожидали_слайдов:
        ошибки.append(f"слайдов в файле {n}, собрано {ожидали_слайдов}")
    for i, s in enumerate(prs.slides, 1):
        for shp in s.shapes:
            if not shp.has_text_frame:
                continue
            for p in shp.text_frame.paragraphs:
                for r in p.runs:
                    if r.text.strip() and r.font.name != F:
                        ошибки.append(f"слайд {i}: run «{r.text[:30]}…» "
                                      f"шрифтом «{r.font.name}», а не Calibri")
    return ошибки
